import os
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path="VorldX_Attendance_System_Complete_Architecture.pptx"):
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    BG_DARK = RGBColor(15, 23, 42)        # #0F172A
    CARD_BG = RGBColor(30, 41, 59)        # #1E293B
    CARD_BG_ALT = RGBColor(24, 33, 47)    # Slightly darker card
    BORDER_COLOR = RGBColor(51, 65, 85)   # #334155
    CYAN_ACCENT = RGBColor(56, 189, 248)  # #38BDF8
    INDIGO_ACCENT = RGBColor(129, 140, 248)# #818CF8
    EMERALD_ACCENT = RGBColor(52, 211, 153)# #34D399
    AMBER_ACCENT = RGBColor(251, 191, 36) # #FBBF24
    ROSE_ACCENT = RGBColor(244, 63, 94)   # #F43F5E
    TEXT_LIGHT = RGBColor(248, 250, 252)  # #F8FAFC
    TEXT_MUTED = RGBColor(148, 163, 184)  # #94A3B8
    TEXT_WHITE = RGBColor(255, 255, 255)

    def set_bg(slide, color=BG_DARK):
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = color
        bg_shape.line.fill.background()
        return bg_shape

    def add_header(slide, category, title, subtitle=None):
        # Category pill/badge
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = CYAN_ACCENT

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.6))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_LIGHT

        # Subtitle
        if subtitle:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(0.35))
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            tf_sub.margin_left = tf_sub.margin_right = tf_sub.margin_top = tf_sub.margin_bottom = 0
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitle
            p_sub.font.size = Pt(11)
            p_sub.font.color.rgb = TEXT_MUTED

    def add_card(slide, left, top, width, height, title, items, tag=None, tag_color=CYAN_ACCENT):
        # Card shape
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = BORDER_COLOR
        shape.line.width = Pt(1)

        # Content Box
        pad = Inches(0.18)
        content_box = slide.shapes.add_textbox(left + pad, top + pad, width - (pad * 2), height - (pad * 2))
        tf = content_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        # Title paragraph
        p_title = tf.paragraphs[0]
        if tag:
            p_title.text = f"[{tag}] {title}"
        else:
            p_title.text = title
        p_title.font.size = Pt(13)
        p_title.font.bold = True
        p_title.font.color.rgb = tag_color
        p_title.space_after = Pt(8)

        # Items
        for item in items:
            p = tf.add_paragraph()
            p.text = f"•  {item}"
            p.font.size = Pt(9.5)
            p.font.color.rgb = TEXT_LIGHT
            p.space_after = Pt(4)

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)

    # Accent decorative box
    acc = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(4.8))
    acc.fill.solid()
    acc.fill.fore_color.rgb = CARD_BG
    acc.line.color.rgb = CYAN_ACCENT
    acc.line.width = Pt(1.5)

    tb = s1.shapes.add_textbox(Inches(1.3), Inches(2.0), Inches(10.7), Inches(3.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "ENTERPRISE TECHNICAL ARCHITECTURE & SOURCE CODE BREAKDOWN"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = CYAN_ACCENT
    p0.space_after = Pt(14)

    p1 = tf.add_paragraph()
    p1.text = "VorldX Employee Attendance System"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(12)

    p2 = tf.add_paragraph()
    p2.text = "Comprehensive Deep Dive into Every Directory, File, Schema, Endpoint & Security Mechanism"
    p2.font.size = Pt(16)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_after = Pt(24)

    p3 = tf.add_paragraph()
    p3.text = "Stack: React 18 (Vite)  •  React Native (Expo SDK 52)  •  Node.js + Express  •  MongoDB Atlas  •  Dynamic Anti-Proxy QR (18s)"
    p3.font.size = Pt(12)
    p3.font.bold = True
    p3.font.color.rgb = EMERALD_ACCENT

    # Footer on Slide 1
    ft = s1.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.733), Inches(0.4))
    p_ft = ft.text_frame.paragraphs[0]
    p_ft.text = "VorldX Attendance Management Suite  |  Confidential & Proprietary Engineering Documentation"
    p_ft.font.size = Pt(10)
    p_ft.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 2: EXECUTIVE SUMMARY & ARCHITECTURAL FOUNDATION
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_header(s2, "System Vision & Core Principles", "VorldX Attendance Platform Overview", "A multi-tier enterprise platform built to eliminate attendance fraud, streamline approvals, and maintain audit trails")

    add_card(s2, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.8), "Anti-Proxy Anti-Replay", [
        "Cryptographically generated 18-second rolling QR codes.",
        "Single-use atomic consumption via MongoDB findOneAndUpdate.",
        "Prevents employees from photographing or sharing codes via chat.",
        "Deterministic check-in with millisecond server timestamping.",
        "Seamless camera scanning via Expo Camera in native mobile app."
    ], "CORE INNOVATION", CYAN_ACCENT)

    add_card(s2, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.8), "Policy & Quota Engine", [
        "Office check-in window: 09:00 AM start, 09:30 AM hard cutoff.",
        "Policy action on late arrival: MARK_LATE, MARK_ABSENT, or HALF_DAY.",
        "Timezone-safe computations strictly pinned to Asia/Kolkata (IST).",
        "Concurrency-safe Work-From-Home (WFH) quota: Max 5 employees/day.",
        "Atomic validation blocks over-allocation across overlapping ranges."
    ], "GOVERNANCE", AMBER_ACCENT)

    add_card(s2, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8), "Full-Stack Separation", [
        "Admin Web Portal: React 18, Vite, React Router 6, Recharts.",
        "Employee Mobile App: React Native, Expo SDK 52, AsyncStorage.",
        "Backend REST API: Node.js, Express, MongoDB Atlas, JWT, bcryptjs.",
        "Immutable Audit Trail: Every correction and approval logged.",
        "Role-Based Access Control: Super Admin, HR Admin, Employee."
    ], "ARCHITECTURE", EMERALD_ACCENT)

    # =========================================================================
    # SLIDE 3: GLOBAL REPOSITORY & FOLDER TAXONOMY
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_header(s3, "Repository Structure", "Top-Level Folder Hierarchy & Workspace Layout", "Clean separation of concerns with four distinct system modules and centralized root configuration")

    add_card(s3, Inches(0.8), Inches(1.8), Inches(2.75), Inches(4.8), "admin-web/", [
        "Enterprise Web Dashboard.",
        "Framework: React 18 + Vite.",
        "Purpose: Super Admin and HR Admin administrative control portal.",
        "Contains 12 distinct routes, charts, policy editor, and live QR kiosk.",
        "Builds with Vite into zero-dependency static production bundle."
    ], "WEB APP", CYAN_ACCENT)

    add_card(s3, Inches(3.8), Inches(1.8), Inches(2.75), Inches(4.8), "backend/", [
        "Core REST API & Business Logic.",
        "Framework: Node.js + Express.",
        "Persistence: MongoDB Atlas via Mongoose ODM schemas.",
        "Contains 11 route files, 9 data models, 3 services, and test suite.",
        "Enforces JWT auth, rate limiting, and 18-second QR generation."
    ], "SERVER API", INDIGO_ACCENT)

    add_card(s3, Inches(6.8), Inches(1.8), Inches(2.75), Inches(4.8), "employee-app/", [
        "Cross-Platform Native App.",
        "Framework: React Native + Expo.",
        "Purpose: Mobile attendance portal for employees & field staff.",
        "Features native Camera QR scanning, WFH/leave requests, and history.",
        "Runs on iOS, Android (APK/Expo Go), and Expo Web."
    ], "MOBILE APP", EMERALD_ACCENT)

    add_card(s3, Inches(9.8), Inches(1.8), Inches(2.75), Inches(4.8), "qr-display/ & Root", [
        "qr-display/: Standalone reception kiosk module for wall tablets.",
        "README.md: Complete engineering setup & architecture guide.",
        "PRESENTATION_SLIDES.md: Coordinator presentation talk-track.",
        "presentation.html: Interactive HTML/CSS slide deck.",
        ".gitignore: Workspace hygiene excluding node_modules & keys."
    ], "SUPPORT & ROOT", AMBER_ACCENT)

    # =========================================================================
    # SLIDE 4: ROOT DIRECTORY FILES IN DETAIL
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_header(s4, "Root Directory Breakdown", "Detailed Examination of Root Files", "Top-level documentation, dependency locks, and presentation files governing the workspace")

    add_card(s4, Inches(0.8), Inches(1.8), Inches(5.7), Inches(2.3), "README.md", [
        "Location: Root directory (`/README.md`) | 100 Lines.",
        "Content: Complete architectural specification and system role breakdown.",
        "Defines default seed credentials (superadmin@vorldx.com / vorldxadminister@gmail.com).",
        "Outlines 09:00 AM start, 09:30 AM cutoff, late policy options, and 5 WFH/day limits.",
        "Provides exact PowerShell execution commands for backend, admin-web, and employee-app."
    ], "DOCUMENTATION", CYAN_ACCENT)

    add_card(s4, Inches(6.8), Inches(1.8), Inches(5.7), Inches(2.3), "PRESENTATION_SLIDES.md", [
        "Location: Root directory (`/PRESENTATION_SLIDES.md`) | 132 Lines.",
        "Content: Slide-by-slide speaker notes and structural script for faculty/coordinator review.",
        "Explains the 18-second rolling QR cryptographic algorithm and atomic queries.",
        "Details Role-Based Access Control matrix (Super Admin vs HR Admin vs Employee).",
        "Highlights automated test verification (api.test.js) and 2,278 Vite modules transformed."
    ], "PRESENTATION", AMBER_ACCENT)

    add_card(s4, Inches(0.8), Inches(4.3), Inches(5.7), Inches(2.3), "presentation.html", [
        "Location: Root directory (`/presentation.html`) | 22.7 KB.",
        "Content: Standalone HTML5/CSS3 presentation deck renderable in any web browser.",
        "Interactive navigation keyboard shortcuts (Space / Arrow keys) for live demos.",
        "Styled with high-contrast corporate dark theme, cards, and data comparison tables.",
        "Requires zero dependencies or local build tools to present project highlights."
    ], "SLIDE DECK UI", EMERALD_ACCENT)

    add_card(s4, Inches(6.8), Inches(4.3), Inches(5.7), Inches(2.3), "package-lock.json & .gitignore", [
        "Location: Root directory (`/package-lock.json`, `/.gitignore`).",
        "package-lock.json: Pins root workspace dependencies and sub-package relationships.",
        ".gitignore: Strictly excludes node_modules, build directories, logs, and sensitive .env files.",
        "Ensures private API keys and database connection strings are never exposed in git.",
        "Maintains consistent reproducibility across Windows, Linux, and macOS environments."
    ], "INFRASTRUCTURE", INDIGO_ACCENT)

    # =========================================================================
    # SLIDE 5: BACKEND OVERVIEW & ARCHITECTURAL STRUCTURE
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_header(s5, "Backend Architecture", "backend/ Directory Overview & Structure", "Modular REST API constructed on Express.js with distinct layers for routing, schemas, and services")

    add_card(s5, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.8), "Directory Structure", [
        "backend/middleware/: Security, JWT verification, and RBAC.",
        "backend/models/: Mongoose ODM schemas and index definitions.",
        "backend/routes/: Modular route handlers (11 discrete route files).",
        "backend/services/: Business logic, QR generation, policy engine.",
        "backend/tests/: Node.js built-in automated test suites.",
        "Root backend scripts: Seeding, password resets, verification."
    ], "FOLDER TAXONOMY", INDIGO_ACCENT)

    add_card(s5, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.8), "Design Patterns", [
        "Modular Routing: Clean route isolation preventing monolithic files.",
        "Service-Oriented Logic: Policy and QR logic separated from controllers.",
        "Atomic Updates: Race-condition-free token and quota manipulation.",
        "Deterministic Timezone Handling: India Standard Time (Asia/Kolkata).",
        "Centralized Error Handling: Uniform JSON response schema."
    ], "ENGINEERING PATTERNS", CYAN_ACCENT)

    add_card(s5, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8), "Security Standard", [
        "Helmet.js integration to enforce HTTP security headers.",
        "Express Rate Limit: 6,000 req/15min (dev) / 600 req/15min (prod).",
        "CORS Whitelist: Restricts origins to configured dashboard and mobile URLs.",
        "12-round bcrypt salted password hashing for all user accounts.",
        "Stateless JWT tokens with sub payload identifying user record."
    ], "SECURITY PROFILE", ROSE_ACCENT)

    # =========================================================================
    # SLIDE 6: BACKEND SERVER ENTRY POINT & CONFIGURATION
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_header(s6, "Backend Core Files", "server.js, .env, and package.json", "Entry point, environment variables, security middleware, and dependency orchestration")

    add_card(s6, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.8), "backend/server.js (5.3 KB | 191 Lines)", [
        "Application Entry Point: Instantiates Express application on port 5000.",
        "Security Middleware: app.disable('x-powered-by'), helmet(), and cors().",
        "CORS Whitelist: Allows 5173, 5174, 8081 (Metro), 19000, 19006, and private LAN IPs.",
        "Rate Limiting: Mounts express-rate-limit over all /api endpoints.",
        "Health Endpoint: GET /api/health checks server status and mongoose.connection.readyState.",
        "Route Mounting: Connects 11 modular route files under /api base path.",
        "Global Error Handler: Handles Mongoose validation errors, CastError, and Mongo error 11000 (duplicate key conflicts).",
        "startServer() Function: Validates MONGODB_URI, JWT_SECRET, and SUPER_ADMIN_SETUP_KEY before listening."
    ], "ENTRY POINT", CYAN_ACCENT)

    add_card(s6, Inches(6.8), Inches(1.8), Inches(5.7), Inches(2.3), "backend/.env & .gitignore", [
        "PORT: Specifies backend port (default 5000).",
        "MONGODB_URI: Cloud connection URI to MongoDB Atlas database cluster.",
        "JWT_SECRET: Minimum 32-character secret key used to sign and verify tokens.",
        "SUPER_ADMIN_SETUP_KEY: Master authorization token for bootstrap seeding.",
        "CLIENT_URL: Permitted origin for CORS headers (http://localhost:5173).",
        "SUPER_ADMIN_EMAIL / PASSWORD: Default credentials for system initialization."
    ], "CONFIGURATION", AMBER_ACCENT)

    add_card(s6, Inches(6.8), Inches(4.3), Inches(5.7), Inches(2.3), "backend/package.json (643 Bytes)", [
        "Type: 'type': 'module' enabling ECMAScript ES6 import/export syntax.",
        "Scripts: 'start': 'node server.js', 'dev': 'nodemon server.js', 'test': 'node --test tests/api.test.js'.",
        "Dependencies: express, mongoose, bcryptjs, jsonwebtoken, cors, helmet, dotenv, express-rate-limit.",
        "Zero heavyweight legacy dependencies; clean modern Node.js 18+ runtime target."
    ], "PACKAGE DEFINITION", INDIGO_ACCENT)

    # =========================================================================
    # SLIDE 7: BACKEND DATA MODELS (PART 1)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7)
    add_header(s7, "Data Architecture", "backend/models/index.js — User, Employee & Attendance", "Comprehensive Mongoose schemas with validation, indexes, and historical revision tracking")

    add_card(s7, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.8), "User & Employee Schemas", [
        "User Schema:",
        "  • name, email (unique, lowercase, trimmed).",
        "  • passwordHash (select: false to prevent exposure).",
        "  • role: ['SUPER_ADMIN', 'HR_ADMIN', 'EMPLOYEE'].",
        "  • active: Boolean (soft deactivation toggle).",
        "Employee Schema:",
        "  • employeeId: String (unique uppercase e.g. EMP001).",
        "  • user: Ref to User document (1-to-1 unique link).",
        "  • department, designation, phone, active."
    ], "SCHEMAS", CYAN_ACCENT)

    add_card(s7, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.8), "Attendance Schema", [
        "employee: ObjectId referencing Employee document.",
        "date: String formatted as YYYY-MM-DD.",
        "mode: Enum ['OFFICE', 'WFH'].",
        "status: Enum ['PRESENT', 'LATE', 'ABSENT', 'HALF_DAY'].",
        "checkIn: Date timestamp (null if absent).",
        "checkOut: Date timestamp (updated on checkout).",
        "checkInMethod: ['QR', 'HR_MANUAL', 'SUPER_ADMIN_MANUAL', 'WFH'].",
        "Unique Compound Index: { employee: 1, date: 1 } ensures strict 1-record-per-day rule."
    ], "CORE ATTENDANCE", EMERALD_ACCENT)

    add_card(s7, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8), "Correction History Sub-Schema", [
        "Embedded sub-document array: correctionHistory: [correctionEntrySchema].",
        "Tracks every administrative manual modification:",
        "  • correctedBy: ObjectId ref to User (Admin identity).",
        "  • correctedAt: Timestamp of modification.",
        "  • reason: Mandatory explanation (max 500 chars).",
        "  • previousStatus vs newStatus.",
        "  • previousCheckIn / previousCheckOut vs new values.",
        "Guarantees 100% auditability; records are never silently changed."
    ], "AUDITABILITY", AMBER_ACCENT)

    # =========================================================================
    # SLIDE 8: BACKEND DATA MODELS (PART 2)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8)
    add_header(s8, "Data Architecture", "backend/models/index.js — Governance & Ephemeral Schemas", "WFH requests, leave management, rolling QR sessions, policy, and audit log models")

    add_card(s8, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.8), "WfhRequest & LeaveRequest", [
        "WfhRequest Model:",
        "  • employee: Ref to Employee.",
        "  • startDate, endDate (YYYY-MM-DD format).",
        "  • status: ['PENDING', 'APPROVED', 'REJECTED', 'CANCELLED'].",
        "  • reason, decisionComment, decidedBy, decidedAt.",
        "LeaveRequest Model:",
        "  • leaveType: ['Casual', 'Sick', 'Earned', 'Emergency'].",
        "  • Compound indexes: { employee: 1, status: 1 } and { startDate: 1, endDate: 1, status: 1 }."
    ], "WORKFLOW MODELS", CYAN_ACCENT)

    add_card(s8, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.8), "QrSession & TTL Index", [
        "token: String (48-char cryptographically secure hex).",
        "createdBy: ObjectId ref to User.",
        "officeLocation: Location string (e.g., 'Main Office').",
        "used: Boolean flag (set atomically upon scan).",
        "usedBy: ObjectId ref to User who scanned.",
        "expiresAt: Timestamp set exactly 18 seconds ahead.",
        "TTL Auto-Purge Index: { expiresAt: 1 } with expireAfterSeconds: 60 cleans up old database documents automatically."
    ], "EPHEMERAL QR", ROSE_ACCENT)

    add_card(s8, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8), "Policy, Audit & Organization", [
        "AttendancePolicy Model:",
        "  • officeStartTime (09:00), checkInCutoffTime (09:30).",
        "  • lateArrivalPolicy: ['MARK_LATE', 'MARK_ABSENT', 'MARK_HALF_DAY'].",
        "  • maxWfhPerDay (default 5), qrExpirySeconds (default 18).",
        "AuditLog Model:",
        "  • actor, actorEmail, actorRole, action, targetType, targetId, details, ipAddress.",
        "Organization Model:",
        "  • Company name, office hours, timezone."
    ], "POLICY & SYSTEM", INDIGO_ACCENT)

    # =========================================================================
    # SLIDE 9: BACKEND SECURITY & MIDDLEWARE
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_bg(s9)
    add_header(s9, "Backend Security Layer", "backend/middleware/auth.js", "Stateless JSON Web Token authentication and declarative Role-Based Access Control guards")

    add_card(s9, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.8), "authenticate Middleware (Lines 4–38)", [
        "Extracts HTTP Authorization header from incoming requests.",
        "Validates header structure: checks for 'Bearer <token>' prefix.",
        "Verifies JWT signature against process.env.JWT_SECRET and issuer 'vorldx-api'.",
        "Decodes subject (payload.sub) and queries User collection in MongoDB.",
        "Validates account active status: if !user or !user.active, returns 401 Unauthorized ('Account is inactive or unavailable').",
        "Attaches authenticated user document to Express request pipeline as req.user.",
        "Catches token expirations and signature tampering, returning standard 401 JSON."
    ], "AUTHENTICATION", CYAN_ACCENT)

    add_card(s9, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), "allowRoles Middleware (Lines 40–50)", [
        "Higher-order middleware accepting allowed role parameters: allowRoles(...roles).",
        "Inspects req.user.role against declared whitelist.",
        "Supported System Roles: 'SUPER_ADMIN', 'HR_ADMIN', 'EMPLOYEE'.",
        "Enforces zero-trust privilege boundaries:",
        "  • If role not included, halts request pipeline immediately.",
        "  • Returns HTTP 403 Forbidden with message: 'You do not have permission to perform this action'.",
        "Applied to routes to prevent privilege escalation (e.g. HR Admin modifying Super Admin)."
    ], "AUTHORIZATION (RBAC)", AMBER_ACCENT)

    # =========================================================================
    # SLIDE 10: BACKEND CORE SERVICES
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_bg(s10)
    add_header(s10, "Backend Services Layer", "policyService.js, qrService.js & auditService.js", "Isolated business logic engines executing policy evaluation, token crypto, and immutable auditing")

    add_card(s10, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.8), "qrService.js (88 Lines)", [
        "createQrSession(userId, officeLocation):",
        "  • Uses crypto.randomBytes(24).toString('hex') to generate 48-character cryptographic token.",
        "  • Sets expiresAt based on dynamic policy (default 18 seconds).",
        "consumeQrToken(token, user):",
        "  • Executes atomic findOneAndUpdate({ token, used: false, expiresAt: { $gt: now } }).",
        "  • Sets used: true, usedBy: user._id, usedAt: now in one atomic step.",
        "  • Eliminates race conditions and replay attacks."
    ], "QR CRYPTO ENGINE", CYAN_ACCENT)

    add_card(s10, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.8), "policyService.js (84 Lines)", [
        "getPolicy() with in-memory caching (60s TTL) for high throughput.",
        "getIndiaDate(date): Uses Intl.DateTimeFormat with Asia/Kolkata timezone returning YYYY-MM-DD.",
        "getIndiaTime(date): Formats 24-hr HH:mm in IST.",
        "evaluateAttendanceStatus(timeStr, policy):",
        "  • Compares check-in time against policy.checkInCutoffTime (09:30).",
        "  • If <= 09:30: Returns 'PRESENT'.",
        "  • If > 09:30: Returns 'LATE', 'ABSENT', or 'HALF_DAY'."
    ], "POLICY ENGINE", EMERALD_ACCENT)

    add_card(s10, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8), "auditService.js (34 Lines)", [
        "logAction({ actor, action, targetType, targetId, details, ipAddress }):",
        "  • Centralized audit logger called across all administrative mutations.",
        "  • Records actor ObjectId, actorEmail, actorRole, action code, target, and IP.",
        "  • Wrapped in resilient try/catch to ensure audit logging never disrupts primary transaction.",
        "  • Provides forensic traceability for compliance audits."
    ], "AUDIT SERVICE", AMBER_ACCENT)

    # =========================================================================
    # SLIDE 11: BACKEND CORE ROUTES (AUTH & ATTENDANCE)
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_bg(s11)
    add_header(s11, "Backend API Routes", "authRoutes.js & attendanceRoutes.js", "Authentication endpoints, mobile check-in/out, punch evaluations, and manual correction history")

    add_card(s11, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.8), "backend/routes/authRoutes.js (6.9 KB)", [
        "POST /api/auth/login: Validates credentials, compares bcrypt hash, generates 7-day JWT.",
        "POST /api/auth/setup-super-admin: Protected bootstrap endpoint requiring SUPER_ADMIN_SETUP_KEY.",
        "GET /api/auth/me: Returns current authenticated user profile, role, and linked employee details.",
        "POST /api/auth/change-password: Allows authenticated users to change passwords with old password verification.",
        "Logs all successful logins and password changes through auditService.js."
    ], "AUTHENTICATION ROUTES", INDIGO_ACCENT)

    add_card(s11, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), "backend/routes/attendanceRoutes.js (15.1 KB | 520 Lines)", [
        "GET /api/attendance/my: Retrieves personal 90-day attendance history for mobile app.",
        "GET /api/attendance/today: Returns today's punch status and checkedIn/checkedOut flags.",
        "POST /api/attendance/check-in: Enforces mode ('OFFICE' vs 'WFH'). Validates 18s QR token or active approved WFH request, evaluates 09:30 IST cutoff, and marks status.",
        "PATCH /api/attendance/check-out: Records official clock-out timestamp.",
        "GET /api/attendance/today-org: Aggregates workforce stats for admin dashboard.",
        "POST /api/attendance/manual & PATCH /api/attendance/:id: Allows Super Admin and HR Admin to correct records, requiring mandatory reason and appending to correctionHistory array."
    ], "ATTENDANCE ROUTES", CYAN_ACCENT)

    # =========================================================================
    # SLIDE 12: BACKEND WORKFLOW ROUTES (WFH, LEAVE & QR)
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_bg(s12)
    add_header(s12, "Backend API Routes", "wfhRoutes.js, leaveRoutes.js & qrRoutes.js", "Work-from-home quota enforcement, leave approval pipelines, and QR session generation")

    add_card(s12, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.8), "backend/routes/wfhRoutes.js (10.1 KB)", [
        "POST /api/wfh: Employees submit WFH requests; validates date ranges and prevents overlaps.",
        "GET /api/wfh/my: Employee's past and active WFH applications.",
        "GET /api/wfh/pending: Queue for Admins to review requests.",
        "PATCH /api/wfh/:id/status: Atomic approval/rejection with quota check. Ensures approved count for any date never exceeds policy.maxWfhPerDay (5).",
        "GET /api/wfh/availability: Returns booked count per date."
    ], "WFH QUOTA ENGINE", AMBER_ACCENT)

    add_card(s12, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.8), "backend/routes/leaveRoutes.js (8.1 KB)", [
        "POST /api/leave: Employees submit leaves with type (Casual, Sick, Earned, Emergency).",
        "Validates date logic and prevents overlapping approved leaves.",
        "GET /api/leave/my: Employee's personal leave ledger.",
        "GET /api/leave/pending: Administrative review queue.",
        "PATCH /api/leave/:id/status: Admin approve/reject endpoint with required decision comments.",
        "Emits audit log entries upon decision."
    ], "LEAVE MANAGEMENT", EMERALD_ACCENT)

    add_card(s12, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8), "backend/routes/qrRoutes.js (1.1 KB)", [
        "POST /api/qr/session: Admin-authenticated endpoint generating dynamic 18-second QR session for office reception.",
        "GET /api/qr/kiosk: Dedicated kiosk display endpoint allowing wall-mounted tablets to continuously fetch fresh tokens without manual re-login.",
        "Integrates directly with qrService.js to output 48-char random tokens and exact expiry timestamps."
    ], "QR ROTATION API", CYAN_ACCENT)

    # =========================================================================
    # SLIDE 13: BACKEND ADMINISTRATIVE & REPORTING ROUTES
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_bg(s13)
    add_header(s13, "Backend Administrative Routes", "employeeRoutes, hrRoutes, auditRoutes, reportRoutes & policyRoutes", "Workforce directory CRUD, HR Admin provisioning, policy updates, and aggregated reporting")

    add_card(s13, Inches(0.8), Inches(1.8), Inches(5.7), Inches(2.3), "employeeRoutes.js & hrRoutes.js", [
        "employeeRoutes.js: Full CRUD for employee records. Enforces unique employee IDs and emails. Allows Super Admin and HR Admin to create, update, or deactivate employees.",
        "hrRoutes.js: Strictly restricted to SUPER_ADMIN role. Allows Super Admin to provision, view, and deactivate HR Admin accounts. HR Admins are forbidden from managing other HR accounts."
    ], "DIRECTORY & HR GOVERNANCE", CYAN_ACCENT)

    add_card(s13, Inches(6.8), Inches(1.8), Inches(5.7), Inches(2.3), "policyRoutes.js & organizationRoutes.js", [
        "policyRoutes.js: GET /api/policy (all admins) and PUT /api/policy (SUPER_ADMIN only). Modifies office start (09:00), cutoff (09:30), late policy, max WFH, and QR validity.",
        "organizationRoutes.js: Configures organizational profile, working hours, and timezones."
    ], "POLICIES & ORG SETTINGS", AMBER_ACCENT)

    add_card(s13, Inches(0.8), Inches(4.3), Inches(5.7), Inches(2.3), "auditRoutes.js (1.3 KB)", [
        "GET /api/audit-logs: Restricted strictly to SUPER_ADMIN.",
        "Returns paginated, searchable timeline of all system actions.",
        "Filterable by actor, action type, date range, and target entity.",
        "Provides cryptographic accountability across all administrative operations."
    ], "AUDIT TRAIL", ROSE_ACCENT)

    add_card(s13, Inches(6.8), Inches(4.3), Inches(5.7), Inches(2.3), "reportRoutes.js (3.7 KB)", [
        "GET /api/reports/summary: Aggregates workforce metrics (total present, late, absent, WFH).",
        "GET /api/reports/department: Groups attendance performance by department.",
        "GET /api/reports/monthly: Aggregates monthly trends for Recharts SVG area graphs.",
        "GET /api/reports/export: Generates structured CSV streams for payroll software import."
    ], "ANALYTICS & EXPORT", EMERALD_ACCENT)

    # =========================================================================
    # SLIDE 14: BACKEND SCRIPTS & TEST SUITE
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    set_bg(s14)
    add_header(s14, "Backend Utility & QA Scripts", "Database Seeding, Recovery & Automated Tests", "Administrative CLI utilities for environment bootstrapping, credential recovery, and test verification")

    add_card(s14, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.8), "Seed & Recovery Scripts", [
        "seedSuperAdmin.js (1.8 KB): Creates initial Super Admin account using environment variables, setting 12-round bcrypt password hash.",
        "seedHRAdmin.js (1.3 KB): Provisions default HR Admin account (vorldxadminister@gmail.com) for day-to-day workforce management.",
        "seedAdmin.js (1.5 KB): Combined bootstrapping script for initial local environment setup.",
        "resetAdminPassword.js (1.2 KB): Standalone CLI recovery utility to reset administrative passwords directly via MongoDB URI.",
        "forceResetSuperAdmin.js (590 Bytes): Emergency script to forcefully restore Super Admin access in case of lockout.",
        "verifyAdmin.js (1.3 KB): Verifies database connection and checks administrative credentials against stored hashes."
    ], "UTILITY & RECOVERY", AMBER_ACCENT)

    add_card(s14, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), "backend/tests/api.test.js (9.6 KB)", [
        "Node.js Native Test Runner: Uses node:test and node:assert.",
        "Automated Test Coverage:",
        "  • Health Check: Verifies /api/health returns 200 OK and database connection.",
        "  • Authentication: Verifies JWT issuance, invalid password rejection, and deactivated user lockout.",
        "  • Role Authorization: Validates that HR Admins cannot access Super Admin endpoints (/api/hr, /api/audit-logs).",
        "  • QR Session Lifecycle: Tests 18-second expiry and verifies that double-consumption throws HTTP 400.",
        "  • WFH Limit: Simulates concurrent requests to verify max 5 approved WFH limit enforcement.",
        "Runs via 'npm test' with zero external test runners required."
    ], "AUTOMATED QA SUITE", EMERALD_ACCENT)

    # =========================================================================
    # SLIDE 15: ADMIN WEB PORTAL OVERVIEW
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    set_bg(s15)
    add_header(s15, "Frontend Architecture", "admin-web/ Directory Overview & Structure", "Modern Single-Page Application (SPA) built with React 18, Vite, and modular component hierarchy")

    add_card(s15, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.8), "Directory Structure", [
        "admin-web/src/components/: Reusable UI primitives.",
        "admin-web/src/context/: Global AuthContext state provider.",
        "admin-web/src/layouts/: AdminLayout wrapper with responsive sidebar.",
        "admin-web/src/pages/: 12 distinct page views.",
        "admin-web/src/services/: Axios API client with interceptors.",
        "admin-web/src/styles/: CSS modules and design system."
    ], "PORTAL TAXONOMY", CYAN_ACCENT)

    add_card(s15, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.8), "Technology Stack", [
        "React 18: Concurrent rendering and performant UI updates.",
        "Vite: Ultra-fast build tool with instant Hot Module Replacement (HMR).",
        "React Router v6: Declarative routing with nested protected routes.",
        "Lucide React: Clean, modern SVG icon library.",
        "Recharts: Composable SVG charting library for attendance trends.",
        "Axios: Promise-based HTTP client."
    ], "TECH HIGHLIGHTS", INDIGO_ACCENT)

    add_card(s15, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8), "Build & Performance", [
        "Optimized production build with zero warnings.",
        "2,278 modules transformed in ~1.4 seconds by Vite.",
        "Modular CSS with CSS custom variables for seamless theme consistency.",
        "Mobile-responsive layout with collapsible drawer navigation.",
        "Stateless token storage in localStorage with auto-logout on 401."
    ], "PERFORMANCE METRICS", EMERALD_ACCENT)

    # =========================================================================
    # SLIDE 16: ADMIN WEB CORE CONFIGURATION & ENTRY POINT
    # =========================================================================
    s16 = prs.slides.add_slide(blank_layout)
    set_bg(s16)
    add_header(s16, "Admin Web Core Files", "package.json, vite.config.js, index.html, main.jsx & App.jsx", "Client-side routing table, role guard orchestration, and build configurations")

    add_card(s16, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.8), "Routing Architecture (App.jsx | 77 Lines)", [
        "LandingRedirect Component: Directs authenticated users to /dashboard or /unauthorized based on role.",
        "Public Routes: / (LoginPage), /unauthorized (Unauthorized access screen).",
        "Protected Shared Admin Routes (Super Admin & HR Admin):",
        "  • /dashboard: Executive overview & metrics.",
        "  • /employees: Staff directory management.",
        "  • /attendance: Daily attendance sheet & corrections.",
        "  • /approvals: WFH and leave approval queue.",
        "  • /office-qr: Dedicated live 18s QR kiosk.",
        "  • /reports: Analytics & CSV export.",
        "  • /settings: Policy view & configuration.",
        "Super Admin Exclusive Nested Routes:",
        "  • /hr-admins: HR credential provisioning.",
        "  • /audit-logs: Immutable system audit logs.",
        "  • /organization: Company profile configuration."
    ], "ROUTING ORCHESTRATION", CYAN_ACCENT)

    add_card(s16, Inches(6.8), Inches(1.8), Inches(5.7), Inches(2.3), "vite.config.js & index.html", [
        "vite.config.js: Configures @vitejs/plugin-react for JSX transformation. Sets dev server port and proxy configurations.",
        "index.html: HTML5 shell with viewport meta tags, Inter/Plus Jakarta Sans Google Fonts, and root DOM mount point (#root)."
    ], "BUILD & DOM SHELL", INDIGO_ACCENT)

    add_card(s16, Inches(6.8), Inches(4.3), Inches(5.7), Inches(2.3), "main.jsx & package.json", [
        "main.jsx: Application bootstrap mounting App within BrowserRouter and AuthProvider using React 18 createRoot API.",
        "package.json: Declares scripts (dev, build, preview) and dependencies (axios, lucide-react, recharts, react-router-dom)."
    ], "BOOTSTRAP & PACKAGES", AMBER_ACCENT)

    # =========================================================================
    # SLIDE 17: ADMIN WEB SERVICES & CONTEXT
    # =========================================================================
    s17 = prs.slides.add_slide(blank_layout)
    set_bg(s17)
    add_header(s17, "Admin Web Infrastructure", "api.js & AuthContext.jsx", "Centralized API communication client and global authentication state management")

    add_card(s17, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.8), "src/services/api.js (2.0 KB)", [
        "Axios Instance Creation: Sets baseURL to VITE_API_URL or http://localhost:5000/api.",
        "Request Interceptor: Automatically inspects localStorage for 'vorldx_token' and attaches 'Authorization: Bearer <token>' to every outgoing HTTP request.",
        "Response Interceptor: Monitors responses for HTTP 401 Unauthorized status.",
        "Auto-Logout on Expiry: If 401 is received, automatically purges expired token and user object from localStorage and redirects user to login page.",
        "Defines helper API functions for attendance, employees, WFH, leaves, reports, policies, and QR sessions."
    ], "API CLIENT LAYER", CYAN_ACCENT)

    add_card(s17, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), "src/context/AuthContext.jsx (3.6 KB)", [
        "AuthContext & useAuth() Hook: Provides global access to user state across all components.",
        "State Variables: user, token, loading, isAuthenticated.",
        "Initialization: On mount, validates cached token by calling /auth/me to fetch fresh user profile.",
        "login(email, password): Executes authentication request, saves token to localStorage, and updates reactive state.",
        "logout(): Flushes localStorage, clears auth context, and triggers redirect.",
        "Role Helpers: Exports isSuperAdmin and isHrAdmin convenience flags for conditional UI rendering."
    ], "STATE MANAGEMENT", EMERALD_ACCENT)

    # =========================================================================
    # SLIDE 18: ADMIN WEB LAYOUT & REUSABLE UI COMPONENTS
    # =========================================================================
    s18 = prs.slides.add_slide(blank_layout)
    set_bg(s18)
    add_header(s18, "Admin Web UI Primitives", "AdminLayout and Component Library", "Modular, reusable UI components ensuring design consistency across all 12 portal pages")

    add_card(s18, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.8), "AdminLayout.jsx (7.2 KB)", [
        "Sidebar Navigation: Renders brand badge, dynamic navigation links with active state indicator.",
        "Role-Aware Filtering: Automatically hides Super Admin tabs from HR Admin users.",
        "Top Header: Displays current page title, current date in IST, and logged-in user profile pill with role badge.",
        "Responsive Drawer: Collapsible hamburger menu on mobile and tablet screens.",
        "Outlet Container: Embeds child routes with smooth fade transitions."
    ], "PAGE SHELL", CYAN_ACCENT)

    add_card(s18, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.8), "Components: ProtectedRoute & Brand", [
        "ProtectedRoute.jsx (741 Bytes):",
        "  • Validates user authentication status.",
        "  • Verifies allowedRoles prop; redirects unauthorized users to /unauthorized.",
        "Brand.jsx (1.1 KB):",
        "  • Renders SVG VorldX attendance logo icon.",
        "  • Displays corporate typography and system version badge.",
        "  • Consistent header branding."
    ], "GUARDS & BRANDING", INDIGO_ACCENT)

    add_card(s18, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8), "StatCard, StatusPill & Modal", [
        "StatCard.jsx (379 Bytes): Reusable metric display with title, value, icon, and trend indicator.",
        "StatusPill.jsx (547 Bytes): Color-coded badges for statuses (PRESENT, LATE, ABSENT, WFH, PENDING, APPROVED).",
        "Modal.jsx (567 Bytes): Accessible pop-up dialog container with backdrop click dismissal and keyboard ESC handling.",
        "EmptyState.jsx (307 Bytes): Placeholder graphic and message for empty lists."
    ], "UI ATOMS", AMBER_ACCENT)

    # =========================================================================
    # SLIDE 19: ADMIN WEB PAGES (PRIMARY OPERATIONS)
    # =========================================================================
    s19 = prs.slides.add_slide(blank_layout)
    set_bg(s19)
    add_header(s19, "Admin Web Portal Pages", "LoginPage, Dashboard, Employees & Attendance", "Operational workflows for authentication, executive monitoring, directory, and master attendance sheet")

    add_card(s19, Inches(0.8), Inches(1.8), Inches(5.7), Inches(2.3), "LoginPage.jsx (6.9 KB)", [
        "Authentication UI with animated background gradients and glassmorphism styling.",
        "Form validation for corporate email and password with error feedback banners.",
        "Quick demo credential buttons for Super Admin and HR Admin testing.",
        "Disables submit button and shows loading spinner during authentication."
    ], "LOGIN PAGE", CYAN_ACCENT)

    add_card(s19, Inches(6.8), Inches(1.8), Inches(5.7), Inches(2.3), "Dashboard.jsx (7.7 KB)", [
        "Executive Cockpit: 4 StatCards (Total Employees, Present Today, On Leave/WFH, Late Arrivals).",
        "7-Day Attendance Trend: Recharts AreaChart showing daily present vs late curves.",
        "Quick Actions Bar: Fast navigation to Add Employee, Mark Attendance, or Open QR Kiosk.",
        "Live Punch Stream: Real-time table of recent clock-ins with status pills."
    ], "EXECUTIVE DASHBOARD", EMERALD_ACCENT)

    add_card(s19, Inches(0.8), Inches(4.3), Inches(5.7), Inches(2.3), "Employees.jsx (6.5 KB)", [
        "Staff Directory: Search by name, ID, or department with instant client filtering.",
        "Add Employee Modal: Form to register employee, create User account, and set department.",
        "Edit & Deactivate: Toggle employee active status; soft deactivation prevents login.",
        "CSV Export: One-click export of employee roster."
    ], "EMPLOYEE DIRECTORY", AMBER_ACCENT)

    add_card(s19, Inches(6.8), Inches(4.3), Inches(5.7), Inches(2.3), "Attendance.jsx (17.8 KB | Largest Page)", [
        "Master Attendance Ledger: Filterable by date, department, and attendance status.",
        "Manual Punch Marking: Admins can mark attendance with check-in/out timestamps.",
        "Correction Modal: Requires mandatory reason when modifying existing punch.",
        "Audit History Drawer: Displays full correctionHistory array showing who changed what and when."
    ], "ATTENDANCE MASTER", INDIGO_ACCENT)

    # =========================================================================
    # SLIDE 20: ADMIN WEB PAGES (WORKFLOWS & GOVERNANCE)
    # =========================================================================
    s20 = prs.slides.add_slide(blank_layout)
    set_bg(s20)
    add_header(s20, "Admin Web Portal Pages", "Approvals, OfficeQR, Reports, HRAdmins, AuditLogs & Settings", "Approval workflows, dynamic reception kiosk, analytics charts, and security governance")

    add_card(s20, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.8), "Approvals.jsx & OfficeQR.jsx", [
        "Approvals.jsx (13.3 KB):",
        "  • Tabbed interface for WFH and Leave requests.",
        "  • WFH Capacity Indicator: Live bar showing approved count vs policy limit (e.g. 3/5 Booked).",
        "  • Approve / Reject action modals with decision comment fields.",
        "OfficeQR.jsx (8.1 KB):",
        "  • Dedicated reception screen fetching 18-second dynamic QR code.",
        "  • Animated circular timer countdown showing seconds remaining.",
        "  • Fullscreen kiosk mode for office entrance tablets."
    ], "WORKFLOW & KIOSK", CYAN_ACCENT)

    add_card(s20, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.8), "Reports.jsx (8.7 KB)", [
        "Workforce Performance Analytics:",
        "  • Breakdown of Present, Late, Absent, and WFH distributions.",
        "  • Recharts BarChart comparing department punctuality.",
        "  • Monthly presence trend area chart.",
        "  • Date range picker for historical analysis.",
        "  • Structured CSV export for payroll processing."
    ], "ANALYTICS & EXPORT", EMERALD_ACCENT)

    add_card(s20, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8), "HRAdmins, AuditLogs & Settings", [
        "HRAdmins.jsx (4.2 KB): Super Admin screen to provision and deactivate HR Admin accounts.",
        "AuditLogs.jsx (7.6 KB): Searchable immutable audit trail showing actor, action, IP, and timestamp.",
        "Settings.jsx (7.9 KB): Configure office start (09:00), cutoff (09:30), late policy, max WFH, and QR validity.",
        "Organization.jsx (3.6 KB): Company info editor.",
        "Unauthorized.jsx (428 Bytes): 403 error fallback."
    ], "SUPER ADMIN CONTROLS", ROSE_ACCENT)

    # =========================================================================
    # SLIDE 21: ADMIN WEB DESIGN SYSTEM & STYLES
    # =========================================================================
    s21 = prs.slides.add_slide(blank_layout)
    set_bg(s21)
    add_header(s21, "Admin Web Styling", "Design System & CSS Architecture", "Custom CSS stylesheets delivering responsive enterprise dark-mode aesthetics")

    add_card(s21, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.8), "global.css (17.4 KB)", [
        "Core Design Tokens & CSS Variables:",
        "  • --bg-primary: #0b132b, --bg-card: #1c2541.",
        "  • --accent-blue: #3a86ff, --accent-cyan: #00f5d4.",
        "  • --status-present: #10b981, --status-late: #f59e0b.",
        "  • --status-absent: #ef4444, --status-wfh: #6366f1.",
        "Typography: System font stack with smooth font smoothing.",
        "Form controls, button states, inputs, tables, scrollbars, and utility classes."
    ], "GLOBAL TOKENS", CYAN_ACCENT)

    add_card(s21, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.8), "layout.css (5.4 KB)", [
        "Layout Structure:",
        "  • CSS Grid and Flexbox layouts for sidebar and main content.",
        "  • Sidebar fixed navigation with active link indicators.",
        "  • Top navigation bar with sticky positioning.",
        "  • Responsive breakpoints for tablet and mobile screens.",
        "  • Collapsible drawer transitions and backdrop overlays."
    ], "LAYOUT & GRID", INDIGO_ACCENT)

    add_card(s21, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8), "login.css (12.2 KB)", [
        "Dedicated Login Screen Styles:",
        "  • Centered glassmorphic card container with backdrop blur.",
        "  • Animated radial gradient background effects.",
        "  • Brand badge animations and pulse effects.",
        "  • Floating input labels and focus ring outlines.",
        "  • Mobile-optimized touch-friendly inputs."
    ], "LOGIN STYLES", AMBER_ACCENT)

    # =========================================================================
    # SLIDE 22: EMPLOYEE MOBILE APP OVERVIEW
    # =========================================================================
    s22 = prs.slides.add_slide(blank_layout)
    set_bg(s22)
    add_header(s22, "Mobile App Architecture", "employee-app/ Directory Overview & Structure", "Cross-platform mobile application built with React Native and Expo SDK 52 for employees")

    add_card(s22, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.8), "Directory Structure", [
        "employee-app/src/screens/: 6 native mobile screens.",
        "employee-app/src/context/: Mobile AuthContext using AsyncStorage.",
        "employee-app/src/services/: Axios mobile API client.",
        "employee-app/assets/: App icon and splash graphics.",
        "employee-app/App.js: Navigation router and auth state wrapper.",
        "employee-app/app.json: Expo configuration and camera permissions."
    ], "MOBILE TAXONOMY", CYAN_ACCENT)

    add_card(s22, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.8), "Core Features", [
        "Expo Camera Integration: Scans 18-second rolling QR codes directly from office reception display.",
        "Today's Attendance Status: Live visual card showing checkedIn/checkedOut status and arrival time.",
        "WFH Application: Mobile form with date picker and real-time status feedback.",
        "Leave Application: Mobile leave requests across all categories.",
        "90-Day Punch History: Personal calendar and attendance log."
    ], "MOBILE CAPABILITIES", EMERALD_ACCENT)

    add_card(s22, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8), "Expo SDK 52 Platform", [
        "Expo SDK 52: Modern React Native runtime supporting new architecture.",
        "Cross-Platform: Single JavaScript codebase runs on Android and iOS.",
        "Expo Camera: Uses modern CameraView API for high-speed barcode scanning.",
        "AsyncStorage: Persistent offline storage for JWT session tokens.",
        "Safe Area Insets: Adapts to notches, dynamic islands, and system bars."
    ], "RUNTIME PLATFORM", INDIGO_ACCENT)

    # =========================================================================
    # SLIDE 23: EMPLOYEE MOBILE CORE & DASHBOARD SCREEN
    # =========================================================================
    s23 = prs.slides.add_slide(blank_layout)
    set_bg(s23)
    add_header(s23, "Mobile Core & Scanner", "App.js, AuthContext.js, api.js & DashboardScreen.js", "Mobile app shell, persistent token storage, camera QR scanning, and punch logic")

    add_card(s23, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.8), "App.js & Infrastructure", [
        "App.js (1.7 KB | 60 Lines):",
        "  • MainNavigator: Manages active screen state ('dashboard', 'history', 'wfh', 'leave', 'profile').",
        "  • Conditionally renders LoginScreen if unauthenticated, or DashboardScreen with back navigation.",
        "AuthContext.js (2.6 KB):",
        "  • Wraps AsyncStorage to persist 'userToken' and 'userData' across app restarts.",
        "  • Exports login(), logout(), user, token, and loading state.",
        "api.js (3.6 KB):",
        "  • Configures base URL targeting backend API.",
        "  • Interceptor automatically injects Bearer token from AsyncStorage into all requests."
    ], "APP CORE & STORAGE", CYAN_ACCENT)

    add_card(s23, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), "DashboardScreen.js (17.4 KB)", [
        "Central Employee Mobile Hub:",
        "  • Header: Welcome greeting with employee name and formatted IST date.",
        "  • Today's Status Card: Shows status pill (PRESENT, LATE, ABSENT) and checkIn/checkOut timestamps.",
        "  • Integrated Expo Camera Scanner: Opens viewfinder modal to scan office reception QR code.",
        "  • Check-In Logic: Sends scanned token to POST /api/attendance/check-in with mode='OFFICE'.",
        "  • WFH Check-In: Allows instant WFH punch if approved request exists for today.",
        "  • Check-Out Button: Sends PATCH /api/attendance/check-out to close daily punch.",
        "  • Navigation Cards: Quick links to History, Apply WFH, Apply Leave, and Profile."
    ], "DASHBOARD & CAMERA SCANNER", EMERALD_ACCENT)

    # =========================================================================
    # SLIDE 24: EMPLOYEE MOBILE WORKFLOW SCREENS
    # =========================================================================
    s24 = prs.slides.add_slide(blank_layout)
    set_bg(s24)
    add_header(s24, "Mobile Workflow Screens", "HistoryScreen, WfhScreen, LeaveScreen & ProfileScreen", "Employee self-service portals for punch ledger, remote work, leave requests, and profile")

    add_card(s24, Inches(0.8), Inches(1.8), Inches(5.7), Inches(2.3), "HistoryScreen.js (6.8 KB)", [
        "90-Day Attendance Ledger: Scrollable list of past punches.",
        "Displays date, mode badge (OFFICE / WFH), status pill, and exact checkIn/checkOut times.",
        "Summary Metrics: Displays 30-day attendance percentage and punctuality score.",
        "Pull-to-refresh integration to fetch updated attendance data."
    ], "PUNCH HISTORY", CYAN_ACCENT)

    add_card(s24, Inches(6.8), Inches(1.8), Inches(5.7), Inches(2.3), "WfhScreen.js (12.1 KB)", [
        "WFH Application Form: Start date, end date, and reason input.",
        "Validates date selections and prevents weekend misconfigurations.",
        "Pending & Past Requests List: Shows status badges (PENDING, APPROVED, REJECTED).",
        "Displays Admin Decision Comments explaining why requests were approved or rejected."
    ], "WFH PORTAL", AMBER_ACCENT)

    add_card(s24, Inches(0.8), Inches(4.3), Inches(5.7), Inches(2.3), "LeaveScreen.js (13.4 KB)", [
        "Leave Application Interface: Selects category (Casual, Sick, Earned, Emergency).",
        "Date range selector with duration calculation (e.g. 3 Days).",
        "Reason text field and emergency contact details.",
        "Request History: View real-time approval status and reviewer notes."
    ], "LEAVE REQUESTS", INDIGO_ACCENT)

    add_card(s24, Inches(6.8), Inches(4.3), Inches(5.7), Inches(2.3), "ProfileScreen.js (7.2 KB)", [
        "Employee Identity Card: Name, Employee ID, Email, Department, Designation.",
        "App Information: Displays current Expo build version and API endpoint connection.",
        "Security: Link to change password modal.",
        "Sign Out Button: Calls AuthContext.logout() to clear session and return to LoginScreen."
    ], "PROFILE & SETTINGS", ROSE_ACCENT)

    # =========================================================================
    # SLIDE 25: OFFICE QR KIOSK DISPLAY (qr-display/)
    # =========================================================================
    s25 = prs.slides.add_slide(blank_layout)
    set_bg(s25)
    add_header(s25, "Office Reception Kiosk", "qr-display/ Module & Hardware Deployment", "Standalone, lightweight display application designed for wall tablets and reception monitors")

    add_card(s25, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.8), "qr-display/ Overview", [
        "Directory: qr-display/",
        "Purpose: Dedicated standalone reception display for office entrances.",
        "Allows office administrators to mount an iPad, Android tablet, or TV monitor at the office door.",
        "Operates independently without exposing administrative dashboard controls.",
        "Communicates with backend via GET /api/qr/kiosk."
    ], "KIOSK ARCHITECTURE", CYAN_ACCENT)

    add_card(s25, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.8), "Files Breakdown", [
        "qr-display/package.json (428 Bytes):",
        "  • Configured with lightweight frontend tooling.",
        "  • Minimal footprint for fast cold-starts on low-power display tablets.",
        "qr-display/README.md (359 Bytes):",
        "  • Hardware mounting instructions.",
        "  • Kiosk screen setup and auto-refresh configuration.",
        "qr-display/src/:",
        "  • Kiosk display rendering logic."
    ], "MODULE FILES", INDIGO_ACCENT)

    add_card(s25, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8), "Kiosk Display Lifecycle", [
        "1. Kiosk display polls GET /api/qr/kiosk upon boot.",
        "2. Backend creates a 48-char random token expiring in 18 seconds.",
        "3. Kiosk renders high-contrast SVG QR code on black background.",
        "4. Circular animated progress bar counts down from 18s to 0s.",
        "5. Exactly at 0s, kiosk auto-requests next token seamlessly.",
        "6. Employees scan code as they walk through office entrance."
    ], "DISPLAY FLOW", EMERALD_ACCENT)

    # =========================================================================
    # SLIDE 26: 18-SECOND ANTI-PROXY QR PROTOCOL
    # =========================================================================
    s26 = prs.slides.add_slide(blank_layout)
    set_bg(s26)
    add_header(s26, "Security Innovation", "18-Second Anti-Proxy Dynamic QR Protocol", "End-to-end cryptographic sequence eliminating proxy attendance and replay vulnerabilities")

    add_card(s26, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.8), "The Proxy Clock-in Vulnerability", [
        "Traditional Flaws:",
        "  • Static Printed QR: Employees take a photo of the printed code and WhatsApp it to absent colleagues.",
        "  • RFID Swipe Cards: Employees physically hand cards to colleagues to swipe on their behalf.",
        "  • Manual Registers: Signatures are forged with zero timestamp verification.",
        "VorldX Innovation:",
        "  • QR code changes every 18 seconds cryptographically.",
        "  • Single-Use Consumption: Once scanned by one employee, token is permanently marked used: true.",
        "  • Network Latency Window: By the time an employee takes a photo and messages it to a friend, the 18-second window has expired.",
        "  • Result: Physical presence at the office kiosk is strictly required."
    ], "THREAT MODEL", ROSE_ACCENT)

    add_card(s26, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), "Atomic Verification Flow", [
        "Step 1 (Generation): Reception screen requests token. Server creates 48-char hex with expiresAt = now + 18s.",
        "Step 2 (Capture): Employee opens authenticated mobile app camera and scans code.",
        "Step 3 (Atomic Query): Mobile sends POST /api/attendance/check-in with { mode: 'OFFICE', qrToken }.",
        "Step 4 (Database Atomicity): Server executes:",
        "    QrSession.findOneAndUpdate(",
        "      { token, used: false, expiresAt: { $gt: now } },",
        "      { $set: { used: true, usedBy: user._id, usedAt: now } }",
        "    )",
        "Step 5 (Validation): If session is null, returns 400 with exact reason (Expired / Already Used / Invalid).",
        "Step 6 (Attendance Write): Evaluates 09:30 cutoff and commits attendance."
    ], "ATOMIC PROTOCOL", CYAN_ACCENT)

    # =========================================================================
    # SLIDE 27: ROLE-BASED ACCESS CONTROL (RBAC) MATRIX
    # =========================================================================
    s27 = prs.slides.add_slide(blank_layout)
    set_bg(s27)
    add_header(s27, "Governance & Security", "Role-Based Access Control (RBAC) Matrix", "Strict privilege boundaries enforcing separation of duties across all user classes")

    # We will build a neat table on slide 27
    table_shape = s27.shapes.add_table(10, 4, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
    table = table_shape.table
    table.columns[0].width = Inches(5.133)
    table.columns[1].width = Inches(2.2)
    table.columns[2].width = Inches(2.2)
    table.columns[3].width = Inches(2.2)

    headers = ["Feature / System Capability", "Super Admin", "HR Admin", "Employee"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = CYAN_ACCENT
            p.alignment = PP_ALIGN.LEFT if i == 0 else PP_ALIGN.CENTER

    rbac_data = [
        ("Mobile App (Camera QR Scan & Check-in/out)", "No Access", "No Access", "Full Access (Mobile)"),
        ("Apply for WFH & Leaves", "No Access", "No Access", "Full Access"),
        ("View Master Workforce Attendance Directory", "Full Access", "Full Access (Employees)", "Own Record Only"),
        ("Manual Attendance Marking & Corrections", "Full Access (With Audit)", "Full Access (Employees)", "Forbidden"),
        ("Approve / Reject WFH & Leaves (Max 5/day Quota)", "Full Access", "Full Access", "Forbidden"),
        ("Create & Deactivate Employee Accounts", "Full Access", "Full Access", "Forbidden"),
        ("Manage HR Admin Accounts (Provisioning)", "Full Access", "Forbidden", "Forbidden"),
        ("View Forensic System Audit Logs", "Full Access", "Forbidden", "Forbidden"),
        ("Configure Attendance Policies & Cutoffs", "Full Access", "Read-Only", "Forbidden"),
    ]

    for row_idx, data in enumerate(rbac_data, start=1):
        for col_idx, text in enumerate(data):
            cell = table.cell(row_idx, col_idx)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG_ALT if row_idx % 2 == 0 else CARD_BG
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(9.5)
                p.alignment = PP_ALIGN.LEFT if col_idx == 0 else PP_ALIGN.CENTER
                if "Full Access" in text:
                    p.font.color.rgb = EMERALD_ACCENT
                elif "Forbidden" in text or "No Access" in text:
                    p.font.color.rgb = ROSE_ACCENT
                else:
                    p.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 28: ATTENDANCE POLICY & TIMEZONE ENGINE
    # =========================================================================
    s28 = prs.slides.add_slide(blank_layout)
    set_bg(s28)
    add_header(s28, "Policy Engine Deep Dive", "Cutoff Logic, Timezones & WFH Concurrency", "Business rules engine governing office punctuality, cutoff enforcement, and remote quota")

    add_card(s28, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.8), "09:30 AM Cutoff Engine", [
        "Office Start Time: Configured to 09:00 AM.",
        "Grace Period Cutoff: Configured to 09:30 AM.",
        "Arrival Evaluation:",
        "  • Check-in time <= 09:30: Automatically classified as PRESENT.",
        "  • Check-in time > 09:30: Triggers configured company policy action.",
        "Policy Action Configurations:",
        "  1. MARK_LATE: Status recorded as LATE (Default).",
        "  2. MARK_ABSENT: Treated as ABSENT despite physical arrival.",
        "  3. MARK_HALF_DAY: Counts as 0.5 working day."
    ], "PUNCTUALITY LOGIC", AMBER_ACCENT)

    add_card(s28, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.8), "Deterministic Timezones", [
        "Timezone: Asia/Kolkata (Indian Standard Time, UTC+05:30).",
        "Cloud Hosting Problem: Cloud servers (AWS, Render, Heroku) run in UTC (00:00). A check-in at 09:15 AM IST evaluates as 03:45 AM UTC.",
        "VorldX Solution: Pinned Intl.DateTimeFormat with Asia/Kolkata timezone in policyService.js.",
        "Calculates getIndiaDate() and getIndiaTime() identically across all servers regardless of geographical cloud region."
    ], "TIMEZONE SAFETY", CYAN_ACCENT)

    add_card(s28, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8), "WFH Concurrency Safety", [
        "Company Quota: Hard ceiling of maximum 5 employees on WFH per calendar day.",
        "Concurrency Race Condition: If 2 admins simultaneously approve 2 pending requests when 4 are already approved, naive queries would result in 6 approved.",
        "VorldX Atomic Validation: Counts active approved requests across target range inside a single transaction before transitioning status from PENDING to APPROVED."
    ], "CONCURRENCY SAFETY", EMERALD_ACCENT)

    # =========================================================================
    # SLIDE 29: SYSTEM DEPLOYMENT & VERIFICATION
    # =========================================================================
    s29 = prs.slides.add_slide(blank_layout)
    set_bg(s29)
    add_header(s29, "Operations & Verification", "Deployment Architecture & Production Build", "Deployment commands, environment initialization, and automated verification status")

    add_card(s29, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.8), "Execution Commands (PowerShell)", [
        "1. Backend Server Initialization:",
        "    cd backend",
        "    npm install",
        "    node seedSuperAdmin.js   # Bootstrap Super Admin",
        "    node seedHRAdmin.js      # Bootstrap HR Admin",
        "    npm test                 # Run automated test suite",
        "    npm start                # Listening on :5000",
        "2. Admin Web Portal:",
        "    cd admin-web",
        "    npm install",
        "    npm run dev              # Vite dev server on :5173",
        "    npm run build            # Production build",
        "3. Employee Mobile App:",
        "    cd employee-app",
        "    npx expo start           # Metro bundler on :8081"
    ], "RUN & DEPLOY COMMANDS", INDIGO_ACCENT)

    add_card(s29, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), "Production Build Verification", [
        "Vite Production Compilation Status:",
        "  • 2,278 modules transformed successfully.",
        "  • Zero compile warnings or syntax errors.",
        "  • Output generated in admin-web/dist/ ready for CDN hosting (Netlify / Vercel / Nginx).",
        "Backend API Health Check:",
        "  • Health endpoint: GET http://localhost:5000/api/health.",
        "  • Database status: MongoDB Atlas connected (readyState: 1).",
        "Mobile App Readiness:",
        "  • Validated on Expo Go and Android Emulator.",
        "  • CameraView permissions configured in app.json."
    ], "QA & COMPILATION STATUS", EMERALD_ACCENT)

    # =========================================================================
    # SLIDE 30: CONCLUSION & FUTURE ROADMAP
    # =========================================================================
    s30 = prs.slides.add_slide(blank_layout)
    set_bg(s30)

    # Final Summary Card
    end_card = s30.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(11.733), Inches(5.5))
    end_card.fill.solid()
    end_card.fill.fore_color.rgb = CARD_BG
    end_card.line.color.rgb = CYAN_ACCENT
    end_card.line.width = Pt(1.5)

    tb_end = s30.shapes.add_textbox(Inches(1.2), Inches(1.5), Inches(10.9), Inches(4.8))
    tf_end = tb_end.text_frame
    tf_end.word_wrap = True

    p_end0 = tf_end.paragraphs[0]
    p_end0.text = "PROJECT SUMMARY & FUTURE INNOVATION ROADMAP"
    p_end0.font.size = Pt(12)
    p_end0.font.bold = True
    p_end0.font.color.rgb = CYAN_ACCENT
    p_end0.space_after = Pt(12)

    p_end1 = tf_end.add_paragraph()
    p_end1.text = "VorldX Attendance System — Fully Implemented & Production-Ready"
    p_end1.font.size = Pt(24)
    p_end1.font.bold = True
    p_end1.font.color.rgb = TEXT_WHITE
    p_end1.space_after = Pt(14)

    p_end2 = tf_end.add_paragraph()
    p_end2.text = "Key Accomplishments:"
    p_end2.font.size = Pt(13)
    p_end2.font.bold = True
    p_end2.font.color.rgb = EMERALD_ACCENT
    p_end2.space_after = Pt(6)

    bullets = [
        "Eliminated Attendance Fraud: 18-second rolling dynamic cryptographic QR codes prevent photo-sharing and proxy check-ins.",
        "Full-Stack Separation: Clean boundary between React 18 Admin Portal, Node.js REST API, and Expo React Native mobile app.",
        "Enterprise Governance: Concurrency-safe WFH quota (max 5/day), 09:30 AM cutoff policy, and immutable correction audit history.",
        "Every File & Folder Documented: 100% architectural coverage across backend, admin-web, employee-app, and qr-display.",
    ]
    for b in bullets:
        pb = tf_end.add_paragraph()
        pb.text = f"✔  {b}"
        pb.font.size = Pt(11)
        pb.font.color.rgb = TEXT_LIGHT
        pb.space_after = Pt(4)

    p_end3 = tf_end.add_paragraph()
    p_end3.text = "\nFuture Roadmap Enhancements:"
    p_end3.font.size = Pt(13)
    p_end3.font.bold = True
    p_end3.font.color.rgb = AMBER_ACCENT
    p_end3.space_after = Pt(6)

    future_bullets = [
        "Geofencing Dual-Check: Verify mobile device GPS coordinates match office latitude/longitude before accepting QR scan.",
        "Face Biometrics: Secondary facial verification during check-in for high-security enterprise installations.",
        "Automated Weekly Slack/Email Digest: Automated notifications to department heads with workforce punctuality analytics."
    ]
    for fb in future_bullets:
        pfb = tf_end.add_paragraph()
        pfb.text = f"➜  {fb}"
        pfb.font.size = Pt(11)
        pfb.font.color.rgb = TEXT_MUTED
        pfb.space_after = Pt(4)

    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    create_deck()
